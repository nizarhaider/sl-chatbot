import io
import os
import uuid
from datetime import datetime, timedelta, timezone

import psycopg
from psycopg.types.json import Json
from fastapi import APIRouter, File, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse
from openpyxl import load_workbook
from pypdf import PdfReader

router = APIRouter(prefix="/app")


def _db():
    value = os.environ.get("DATABASE_URL")
    if not value:
        raise HTTPException(503, "Database is not configured")
    return value


def _customer(connection):
    phone_number_id = os.environ.get("PHONE_NUMBER_ID")
    row = connection.execute("select customer_id from whatsapp_numbers where phone_number_id = %s and status = 'active' limit 1", (phone_number_id,)).fetchone()
    if row is None:
        raise HTTPException(503, "No active customer is configured")
    return str(row[0])


def _schema(connection, customer_id):
    connection.execute("""
        create table if not exists agent_profiles (
            id uuid primary key, customer_id uuid not null, name text not null, instructions text not null,
            enabled_tools jsonb not null default '[]', active boolean not null default true, created_at timestamptz not null default now()
        )
    """)
    connection.execute("""
        create table if not exists knowledge_documents (
            id uuid primary key, customer_id uuid not null, filename text not null, content text not null,
            status text not null default 'indexed', created_at timestamptz not null default now()
        )
    """)
    connection.execute("""
        create table if not exists client_catalog (
            id uuid primary key, customer_id uuid not null, name text not null, description text not null default '',
            price numeric, stock integer not null default 0, status text not null default 'active', updated_at timestamptz not null default now()
        )
    """)
    connection.execute("""
        create table if not exists client_onboarding (
            customer_id uuid primary key, business_name text, display_number text, phone_number_id text,
            meta_status text not null default 'draft', updated_at timestamptz not null default now()
        )
    """)
    count = connection.execute("select count(*) from agent_profiles where customer_id = %s", (customer_id,)).fetchone()[0]
    if not count:
        connection.execute("insert into agent_profiles (id, customer_id, name, instructions, enabled_tools) values (%s,%s,%s,%s,%s)", (uuid.uuid4(), customer_id, "Homelands concierge", "Be a concise, warm female agent. Ask one focused question at a time.", Json(["search_properties", "book_appointment", "send_whatsapp_message", "query_catalog"])) )
    count = connection.execute("select count(*) from client_catalog where customer_id = %s", (customer_id,)).fetchone()[0]
    if not count:
        for name, description, price, stock in [("City Gardens viewing", "Three-bedroom apartment in Rajagiriya.", 41500000, 3), ("Horizon Residencies", "Two-bedroom apartment in Malabe.", 28000000, 0), ("Home consultation", "Property consultation with a Homelands agent.", 0, 12)]:
            connection.execute("insert into client_catalog (id,customer_id,name,description,price,stock) values (%s,%s,%s,%s,%s,%s)", (uuid.uuid4(), customer_id, name, description, price, stock))


def _portal(request):
    if request.cookies.get("serendibai_portal") != "admin":
        raise HTTPException(401, "Sign in required")


def _rows(sql, params=()):
    with psycopg.connect(_db()) as connection:
        customer_id = _customer(connection)
        _schema(connection, customer_id)
        return connection.execute(sql, (customer_id, *params)).fetchall()


@router.post("/login")
async def login(request: Request):
    body = await request.json()
    if body.get("username") != "admin" or body.get("password") != "admin":
        raise HTTPException(401, "Invalid credentials")
    response = JSONResponse({"ok": True})
    response.set_cookie("serendibai_portal", "admin", httponly=True, samesite="lax", secure=True)
    return response


@router.get("", response_class=HTMLResponse)
async def portal():
    return PORTAL_HTML


@router.get("/api/overview")
def overview(request: Request):
    _portal(request)
    with psycopg.connect(_db()) as connection:
        customer_id = _customer(connection)
        _schema(connection, customer_id)
        calls = connection.execute("select count(*), count(*) filter (where status='completed') from calls where customer_id=%s and created_at >= now() - interval '30 days'", (customer_id,)).fetchone()
        appointments = connection.execute("select count(*) from property_appointments where customer_id=%s and status='booked'", (customer_id,)).fetchone()[0]
        catalog = connection.execute("select count(*) from client_catalog where customer_id=%s", (customer_id,)).fetchone()[0]
    return {"calls": calls[0], "completed": calls[1], "appointments": appointments, "catalog": catalog}


@router.get("/api/agents")
def agents(request: Request):
    _portal(request)
    rows = _rows("select id,name,instructions,enabled_tools,active from agent_profiles where customer_id=%s order by created_at")
    return [{"id": str(r[0]), "name": r[1], "instructions": r[2], "enabled_tools": r[3], "active": r[4]} for r in rows]


@router.post("/api/agents")
async def save_agent(request: Request):
    _portal(request)
    data = await request.json()
    enabled_tools = list(data.get("enabled_tools", []))
    if "query_catalog" not in enabled_tools:
        enabled_tools.append("query_catalog")
    with psycopg.connect(_db()) as connection:
        customer_id = _customer(connection)
        _schema(connection, customer_id)
        agent_id = data.get("id") or str(uuid.uuid4())
        connection.execute("insert into agent_profiles (id,customer_id,name,instructions,enabled_tools,active) values (%s,%s,%s,%s,%s,%s) on conflict (id) do update set name=excluded.name,instructions=excluded.instructions,enabled_tools=excluded.enabled_tools,active=excluded.active", (agent_id, customer_id, data.get("name", "Voice agent"), data.get("instructions", ""), Json(enabled_tools), bool(data.get("active", True))))
    return {"ok": True, "id": agent_id}


@router.get("/api/catalog")
def catalog(request: Request):
    _portal(request)
    rows = _rows("select id,name,description,price,stock,status from client_catalog where customer_id=%s order by name")
    return [{"id": str(r[0]), "name": r[1], "description": r[2], "price": float(r[3] or 0), "stock": r[4], "status": r[5]} for r in rows]


@router.post("/api/catalog")
async def save_catalog(request: Request):
    _portal(request)
    data = await request.json()
    with psycopg.connect(_db()) as connection:
        customer_id = _customer(connection)
        _schema(connection, customer_id)
        item_id = data.get("id") or str(uuid.uuid4())
        connection.execute("insert into client_catalog (id,customer_id,name,description,price,stock,status) values (%s,%s,%s,%s,%s,%s,%s) on conflict (id) do update set name=excluded.name,description=excluded.description,price=excluded.price,stock=excluded.stock,status=excluded.status,updated_at=now()", (item_id, customer_id, data.get("name", "Untitled"), data.get("description", ""), data.get("price") or 0, data.get("stock") or 0, data.get("status", "active")))
    return {"ok": True, "id": item_id}


@router.post("/api/catalog/import")
async def import_catalog(request: Request, file: UploadFile = File(...)):
    _portal(request)
    book = load_workbook(io.BytesIO(await file.read()), read_only=True, data_only=True)
    sheet = book.active
    headers = [str(cell.value or "").strip().lower() for cell in next(sheet.iter_rows(max_row=1))]
    required = {"name", "description", "price", "stock"}
    if not required <= set(headers):
        raise HTTPException(400, "Excel needs name, description, price, and stock columns")
    records = [dict(zip(headers, [cell.value for cell in row])) for row in sheet.iter_rows(min_row=2) if row[0].value]
    with psycopg.connect(_db()) as connection:
        customer_id = _customer(connection)
        _schema(connection, customer_id)
        for record in records:
            connection.execute("insert into client_catalog (id,customer_id,name,description,price,stock,status) values (%s,%s,%s,%s,%s,%s,%s)", (uuid.uuid4(), customer_id, record["name"], record.get("description") or "", record.get("price") or 0, record.get("stock") or 0, "active" if (record.get("stock") or 0) else "out_of_stock"))
    return {"ok": True, "imported": len(records)}


@router.get("/api/documents")
def documents(request: Request):
    _portal(request)
    rows = _rows("select id,filename,status,created_at from knowledge_documents where customer_id=%s order by created_at desc")
    return [{"id": str(r[0]), "filename": r[1], "status": r[2], "created_at": r[3].isoformat()} for r in rows]


def _text(filename, content):
    if filename.lower().endswith(".pdf"):
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    return content.decode("utf-8", errors="ignore")


@router.post("/api/documents")
async def upload_document(request: Request, file: UploadFile = File(...)):
    _portal(request)
    content = _text(file.filename or "document", await file.read()).strip()
    if not content:
        raise HTTPException(400, "No readable text found")
    document_id = uuid.uuid4()
    with psycopg.connect(_db()) as connection:
        customer_id = _customer(connection)
        _schema(connection, customer_id)
        connection.execute("insert into knowledge_documents (id,customer_id,filename,content) values (%s,%s,%s,%s)", (document_id, customer_id, file.filename or "document", content))
    _index_document(str(document_id), file.filename or "document", content)
    return {"ok": True, "id": str(document_id), "status": "indexed"}


def _index_document(document_id, filename, content):
    from pinecone import Pinecone
    key = os.environ.get("PINECONE_API_KEY")
    if not key:
        return
    with psycopg.connect(_db()) as connection:
        customer_id = _customer(connection)
    index = Pinecone(api_key=key).Index(os.environ.get("PINECONE_INDEX_NAME", "homelands-properties"))
    records = [{"_id": f"doc:{document_id}:{n}", "content": chunk, "kind": "document", "document_id": document_id, "filename": filename} for n, chunk in enumerate([content[i:i + 1200] for i in range(0, len(content), 1200)])]
    index.upsert_records(namespace=customer_id, records=records)


@router.get("/api/appointments")
def appointments(request: Request):
    _portal(request)
    rows = _rows("select a.id,p.name,a.customer_name,a.customer_phone,a.appointment_at,a.status from property_appointments a join real_estate_properties p on p.id=a.property_id where a.customer_id=%s order by a.appointment_at")
    return [{"id": str(r[0]), "property": r[1], "name": r[2], "phone": r[3], "start": r[4].isoformat(), "status": r[5]} for r in rows]


@router.post("/api/appointments/{appointment_id}")
async def update_appointment(appointment_id: str, request: Request):
    _portal(request)
    data = await request.json()
    with psycopg.connect(_db()) as connection:
        customer_id = _customer(connection)
        connection.execute("update property_appointments set status=%s where id=%s and customer_id=%s", (data.get("status", "booked"), appointment_id, customer_id))
    return {"ok": True}


@router.get("/api/onboarding")
def onboarding(request: Request):
    _portal(request)
    rows = _rows("select business_name,display_number,phone_number_id,meta_status from client_onboarding where customer_id=%s")
    return {} if not rows else {"business_name": rows[0][0], "display_number": rows[0][1], "phone_number_id": rows[0][2], "meta_status": rows[0][3]}


@router.post("/api/onboarding")
async def save_onboarding(request: Request):
    _portal(request)
    data = await request.json()
    with psycopg.connect(_db()) as connection:
        customer_id = _customer(connection)
        _schema(connection, customer_id)
        connection.execute("insert into client_onboarding (customer_id,business_name,display_number,phone_number_id,meta_status) values (%s,%s,%s,%s,%s) on conflict (customer_id) do update set business_name=excluded.business_name,display_number=excluded.display_number,phone_number_id=excluded.phone_number_id,meta_status=excluded.meta_status,updated_at=now()", (customer_id, data.get("business_name"), data.get("display_number"), data.get("phone_number_id"), "pending_meta"))
    return {"ok": True, "status": "pending_meta"}


PORTAL_HTML = """<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>SerendibAI</title><style>:root{font-family:Inter,system-ui;background:#f6f7fb;color:#172033}*{box-sizing:border-box}body{margin:0;display:grid;grid-template-columns:230px 1fr;min-height:100vh}aside{background:#111a2d;color:#fff;padding:24px}h1{font-size:20px;margin:0 0 30px}nav button{display:block;border:0;background:none;color:#b8c2d9;padding:12px 0;width:100%;text-align:left;font:inherit;cursor:pointer}nav button.active{color:#7ce4c6}main{padding:34px;max-width:1250px;width:100%;margin:auto}.page{display:none}.page.active{display:block}.cards{display:grid;grid-template-columns:repeat(4,1fr);gap:14px}.card,section{background:#fff;border:1px solid #e4e8ef;border-radius:14px;padding:20px;margin-bottom:18px}.value{font-size:30px;font-weight:750;margin-top:8px}.muted{color:#718096;font-size:13px}input,textarea,select,button{font:inherit}input,textarea,select{width:100%;border:1px solid #d8deea;border-radius:8px;padding:10px;margin:6px 0 12px}textarea{min-height:140px}button.primary{background:#176b58;color:white;border:0;border-radius:8px;padding:10px 14px;cursor:pointer}table{width:100%;border-collapse:collapse}td,th{text-align:left;padding:11px;border-bottom:1px solid #edf0f5}.pill{padding:4px 8px;border-radius:999px;background:#e8faf4;color:#176b58;font-size:12px}@media(max-width:750px){body{grid-template-columns:1fr}aside{padding:16px}.cards{grid-template-columns:repeat(2,1fr)}} </style></head><body><aside><h1>SerendibAI</h1><div class='muted'>Client control plane</div><nav><button class='active' data-page='dashboard'>Dashboard</button><button data-page='agents'>Agents</button><button data-page='documents'>Documents</button><button data-page='data'>Data</button><button data-page='appointments'>Appointments</button><button data-page='onboarding'>Onboarding</button></nav></aside><main><div id='login'><h2>Sign in</h2><section><input id='user' placeholder='Username' value='admin'><input id='pass' type='password' placeholder='Password' value='admin'><button class='primary' onclick='signin()'>Continue</button></section></div><div id='app' hidden><div id='dashboard' class='page active'><h2>Dashboard</h2><div class='cards' id='metrics'></div><section><h3>Ready for demo</h3><p>Mock agents, catalog items, and appointments are already seeded. Changes here affect the next Live call.</p></section></div><div id='agents' class='page'><h2>Agent management</h2><section><input id='agentName' placeholder='Agent name'><textarea id='agentInstructions'></textarea><label><input type='checkbox' value='search_properties' checked> Property search</label><label><input type='checkbox' value='book_appointment' checked> Booking</label><label><input type='checkbox' value='send_whatsapp_message' checked> WhatsApp messaging</label><button class='primary' onclick='saveAgent()'>Save for next call</button></section></div><div id='documents' class='page'><h2>Documents</h2><section><input type='file' id='doc'><button class='primary' onclick='uploadDoc()'>Upload and index</button><div id='docs'></div></section></div><div id='data' class='page'><h2>Products & services</h2><section><input type='file' id='sheet' accept='.xlsx'><button class='primary' onclick='importSheet()'>Import Excel</button><p class='muted'>Expected columns: name, description, price, stock.</p></section><section><button class='primary' onclick='addItem()'>Add item</button><table id='catalog'></table></section></div><div id='appointments' class='page'><h2>Appointments</h2><section><table id='appointmentsTable'></table></section></div><div id='onboarding' class='page'><h2>WhatsApp onboarding</h2><section><input id='business' placeholder='Business name'><input id='displayNumber' placeholder='WhatsApp display number'><input id='phoneNumberId' placeholder='Meta phone number ID'><button class='primary' onclick='saveOnboarding()'>Save Meta setup</button><p class='muted'>Meta Embedded Signup can be connected here once your Meta app ID and OAuth redirect are supplied.</p></section></div></div></main><script>let agent;const api=(path,o={})=>fetch('/app/api/'+path,{credentials:'same-origin',headers:{'Content-Type':'application/json',...(o.headers||{})},...o}).then(async r=>{if(r.status===401){login.hidden=false;app.hidden=true;throw Error('login')}if(!r.ok)throw Error(await r.text());return r.json()});document.querySelectorAll('nav button').forEach(b=>b.onclick=()=>{document.querySelectorAll('nav button,.page').forEach(x=>x.classList.remove('active'));b.classList.add('active');document.getElementById(b.dataset.page).classList.add('active');load(b.dataset.page)});async function signin(){await api('login',{method:'POST',body:JSON.stringify({username:user.value,password:pass.value})});login.hidden=true;app.hidden=false;load('dashboard')}async function load(page){if(page==='dashboard'){let d=await api('overview');metrics.innerHTML=Object.entries(d).map(([k,v])=>`<div class='card'><div class='muted'>${k}</div><div class='value'>${v}</div></div>`).join('')}if(page==='agents'){agent=(await api('agents'))[0];agentName.value=agent.name;agentInstructions.value=agent.instructions}if(page==='documents'){docs.innerHTML=(await api('documents')).map(d=>`<p>${d.filename} <span class='pill'>${d.status}</span></p>`).join('')}if(page==='data'){let d=await api('catalog');catalog.innerHTML='<tr><th>Name</th><th>Stock</th><th>Status</th></tr>'+d.map(x=>`<tr><td>${x.name}</td><td><input value='${x.stock}' onchange='updateStock("${x.id}",this.value)'></td><td>${x.status}</td></tr>`).join('')}if(page==='appointments'){let d=await api('appointments');appointmentsTable.innerHTML='<tr><th>When</th><th>Property</th><th>Customer</th><th>Status</th></tr>'+d.map(x=>`<tr><td>${new Date(x.start).toLocaleString()}</td><td>${x.property}</td><td>${x.name}</td><td><select onchange='setAppointment("${x.id}",this.value)'><option ${x.status==='booked'?'selected':''}>booked</option><option ${x.status==='cancelled'?'selected':''}>cancelled</option></select></td></tr>`).join('')}if(page==='onboarding'){let d=await api('onboarding');business.value=d.business_name||'';displayNumber.value=d.display_number||'';phoneNumberId.value=d.phone_number_id||''}}async function saveAgent(){let tools=[...document.querySelectorAll('#agents input[type=checkbox]:checked')].map(x=>x.value);await api('agents',{method:'POST',body:JSON.stringify({...agent,name:agentName.value,instructions:agentInstructions.value,enabled_tools:tools,active:true})})}async function uploadDoc(){let f=new FormData();f.append('file',doc.files[0]);await fetch('/app/api/documents',{method:'POST',body:f,credentials:'same-origin'});load('documents')}async function importSheet(){let f=new FormData();f.append('file',sheet.files[0]);await fetch('/app/api/catalog/import',{method:'POST',body:f,credentials:'same-origin'});load('data')}async function updateStock(id,stock){let item=(await api('catalog')).find(x=>x.id===id);await api('catalog',{method:'POST',body:JSON.stringify({...item,stock:+stock,status:stock>0?'active':'out_of_stock'})})}async function addItem(){await api('catalog',{method:'POST',body:JSON.stringify({name:'New product',description:'Demo item',price:0,stock:1,status:'active'})});load('data')}async function setAppointment(id,status){await api('appointments/'+id,{method:'POST',body:JSON.stringify({status})})}async function saveOnboarding(){await api('onboarding',{method:'POST',body:JSON.stringify({business_name:business.value,display_number:displayNumber.value,phone_number_id:phoneNumberId.value})})}login.hidden=false;</script></body></html>"""
